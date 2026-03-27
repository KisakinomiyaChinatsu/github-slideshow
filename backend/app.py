from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
import os
import openpyxl
from openpyxl.utils import get_column_letter
from werkzeug.utils import secure_filename

app = Flask(__name__)
CORS(app)

DATA_DIR = os.path.join(os.path.dirname(__file__), '..', 'data')
os.makedirs(DATA_DIR, exist_ok=True)

ALLOWED_EXTENSIONS = {'xlsx'}

def allowed_file(filename):
    if filename is None:
        return False
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

@app.route('/api/files', methods=['GET'])
def list_files():
    files = []
    for f in os.listdir(DATA_DIR):
        if f.endswith('.xlsx'):
            path = os.path.join(DATA_DIR, f)
            wb = openpyxl.load_workbook(os.path.join(DATA_DIR, f))
            files.append({
                'name': f,
                'sheets': wb.sheetnames,
                'modified': os.path.getmtime(path)
            })
    return jsonify(files)

@app.route('/api/files', methods=['POST'])
def create_file():
    if 'file' not in request.files:
        name = request.json.get('name', 'new_workbook')
        if not name.endswith('.xlsx'):
            name += '.xlsx'
    else:
        file = request.files['file']
        if file.filename == '':
            return jsonify({'error': 'No file selected'}), 400
        name = secure_filename(file.filename)

    path = os.path.join(DATA_DIR, name)
    if os.path.exists(path):
        return jsonify({'error': 'File already exists'}), 400

    wb = openpyxl.Workbook()
    wb.save(path)
    return jsonify({'name': name, 'sheets': wb.sheetnames})

@app.route('/api/files/<name>', methods=['GET'])
def get_file(name):
    path = os.path.join(DATA_DIR, name)
    if not os.path.exists(path):
        return jsonify({'error': 'File not found'}), 404

    wb = openpyxl.load_workbook(path)
    return jsonify({
        'name': name,
        'sheets': wb.sheetnames,
        'activeSheet': wb.active.title
    })

@app.route('/api/files/<name>/sheets/<sheet_name>', methods=['GET'])
def get_sheet(name, sheet_name):
    path = os.path.join(DATA_DIR, name)
    if not os.path.exists(path):
        return jsonify({'error': 'File not found'}), 404

    wb = openpyxl.load_workbook(path)
    sheet = wb[sheet_name]
    if sheet is None:
        return jsonify({'error': 'Sheet not found'}), 404
    data = []
    for row in sheet.iter_rows(values_only=True):
        data.append(list(row))

    return jsonify({
        'name': name,
        'sheet': sheet_name,
        'data': data
    })

@app.route('/api/files/<name>/sheets/<sheet_name>', methods=['PUT'])
def update_sheet(name, sheet_name):
    path = os.path.join(DATA_DIR, name)
    if not os.path.exists(path):
        return jsonify({'error': 'File not found'}), 404

    wb = openpyxl.load_workbook(path)
    if sheet_name not in wb.sheetnames:
        return jsonify({'error': 'Sheet not found'}), 404

    ws = wb[sheet_name]
    new_data = request.json.get('data', [])

    for i, row in enumerate(new_data, start=1):
        for j, value in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=value)

    wb.save(path)
    return jsonify({'success': True})

@app.route('/api/files/<name>/sheets', methods=['POST'])
def create_sheet(name):
    path = os.path.join(DATA_DIR, name)
    if not os.path.exists(path):
        return jsonify({'error': 'File not found'}), 404

    wb = openpyxl.load_workbook(path)
    sheet_name = request.json.get('name', 'Sheet')

    if sheet_name in wb.sheetnames:
        return jsonify({'error': 'Sheet already exists'}), 400

    wb.create_sheet(sheet_name)
    wb.save(path)
    return jsonify({'name': sheet_name, 'sheets': wb.sheetnames})

@app.route('/api/files/<name>', methods=['DELETE'])
def delete_file(name):
    path = os.path.join(DATA_DIR, name)
    if not os.path.exists(path):
        return jsonify({'error': 'File not found'}), 404

    os.remove(path)
    return jsonify({'success': True})

@app.route('/api/files/<name>/download', methods=['GET'])
def download_file(name):
    path = os.path.join(DATA_DIR, name)
    if not os.path.exists(path):
        return jsonify({'error': 'File not found'}), 404
    return send_file(path, as_attachment=True)

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=3001, debug=True)
